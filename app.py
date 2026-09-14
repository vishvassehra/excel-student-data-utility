import streamlit as st
from io import BytesIO
from copy import copy
import openpyxl

st.set_page_config(
    page_title="Excel Student Data Utility",
    page_icon="📊",
    layout="centered"
)

st.title("📊 Excel Student Data Utility")
st.write("Upload your fee collection report. The utility fills blank student details from the row above.")

st.info(
    "Columns processed: A to E — Receipt No., Adm No., Name, Class, Receipt Date. "
    "Existing values are never overwritten."
)

uploaded_file = st.file_uploader(
    "Upload Excel file",
    type=["xlsx", "xlsm"],
    help="Upload the Excel report you want to process."
)

if uploaded_file:
    st.success(f"File selected: {uploaded_file.name}")

    if st.button("⚙️ Process Excel File", type="primary", use_container_width=True):
        try:
            keep_vba = uploaded_file.name.lower().endswith(".xlsm")
            input_bytes = BytesIO(uploaded_file.getvalue())

            wb = openpyxl.load_workbook(input_bytes, keep_vba=keep_vba)
            total_filled = 0

            # Process every worksheet
            for ws in wb.worksheets:
                last_row = ws.max_row

                # Your report starts student data at row 4
                for r in range(4, last_row + 1):
                    for c in range(1, 6):  # A:E
                        current = ws.cell(r, c)
                        above = ws.cell(r - 1, c)

                        is_blank = (
                            current.value is None
                            or (isinstance(current.value, str) and current.value.strip() == "")
                        )

                        if is_blank and above.value is not None:
                            # Copy value
                            current.value = above.value

                            # Preserve number format/alignment/font/border/fill/protection
                            if above.has_style:
                                current._style = copy(above._style)
                            if above.number_format:
                                current.number_format = above.number_format
                            if above.alignment:
                                current.alignment = copy(above.alignment)
                            if above.font:
                                current.font = copy(above.font)
                            if above.fill:
                                current.fill = copy(above.fill)
                            if above.border:
                                current.border = copy(above.border)
                            if above.protection:
                                current.protection = copy(above.protection)

                            total_filled += 1

            output = BytesIO()
            wb.save(output)
            output.seek(0)

            original = uploaded_file.name
            if original.lower().endswith(".xlsm"):
                output_name = original[:-5] + "_filled.xlsm"
            else:
                output_name = original[:-5] + "_filled.xlsx"

            st.success(f"Done! {total_filled} blank cells were filled.")
            st.download_button(
                label="⬇️ Download Processed Excel",
                data=output,
                file_name=output_name,
                mime=(
                    "application/vnd.ms-excel.sheet.macroEnabled.12"
                    if keep_vba
                    else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                ),
                use_container_width=True
            )

        except Exception as e:
            st.error("The file could not be processed.")
            st.exception(e)

st.divider()
st.caption("Only blank cells in columns A:E are filled. Existing data is not changed.")
